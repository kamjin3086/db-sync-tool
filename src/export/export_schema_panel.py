import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading
from typing import Dict, List, Tuple, Optional
from datetime import datetime

import pymysql

from src.components.connection_frame import DbConnectionFrame
from src.components.config_manager import ConfigManager


class ExportSchemaPanel(ttk.Frame):
    """
    将数据库结构导出为 Excel（每个表一个工作表），
    工作表顶部展示表名与表描述，紧随其后为字段清单：
    [字段名, 数据类型（长）, 是否为空, 是否为主键, 说明]
    """

    def __init__(self, master, config_manager: ConfigManager) -> None:
        super().__init__(master, padding="10")
        self.config_manager = config_manager

        # 连接面板（复用）
        self.conn_frame = DbConnectionFrame(self, self.config_manager, text="1. 配置数据库连接")
        self.conn_frame.pack(fill=tk.X, padx=5, pady=5)

        # 保存路径
        save_frame = ttk.LabelFrame(self, text="2. 选择导出文件")
        save_frame.pack(fill=tk.X, padx=5, pady=5)
        self.save_path_var = tk.StringVar()
        ttk.Entry(save_frame, textvariable=self.save_path_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5, pady=5)
        ttk.Button(save_frame, text="浏览...", command=self.browse_save_path).pack(side=tk.LEFT, padx=5, pady=5)

        # 导出设置与表选择
        self.merge_to_one_sheet = tk.BooleanVar(value=False)
        self.tables_data: List[Tuple[str, str]] = []  # [(table, comment)]
        self.table_vars: Dict[str, tk.BooleanVar] = {}

        options_frame = ttk.LabelFrame(self, text="3. 导出设置与表选择（默认全选，可取消）")
        options_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        top_row = ttk.Frame(options_frame)
        top_row.pack(fill=tk.X, padx=5, pady=5)
        ttk.Checkbutton(top_row, text="合并至一张Sheet（按表顺序连续写入）", variable=self.merge_to_one_sheet).pack(side=tk.LEFT, padx=5)

        self.include_toc = tk.BooleanVar(value=True)
        ttk.Checkbutton(top_row, text="生成索引/目录（首个Sheet，内链跳转）", variable=self.include_toc).pack(side=tk.LEFT, padx=5)
        ttk.Button(top_row, text="加载/刷新表列表", command=self.load_tables_list).pack(side=tk.LEFT, padx=5)
        self.toggle_select_btn = ttk.Button(top_row, text="全选", command=self.toggle_select_tables)
        self.toggle_select_btn.pack(side=tk.LEFT, padx=5)

        # 滚动表清单
        self.tables_canvas = tk.Canvas(options_frame, height=200)
        self.tables_scrollbar = ttk.Scrollbar(options_frame, orient="vertical", command=self.tables_canvas.yview)
        self.tables_inner = ttk.Frame(self.tables_canvas)
        self.tables_inner.bind(
            "<Configure>", lambda e: self.tables_canvas.configure(scrollregion=self.tables_canvas.bbox("all"))
        )
        self.tables_canvas.create_window((0, 0), window=self.tables_inner, anchor="nw")
        self.tables_canvas.configure(yscrollcommand=self.tables_scrollbar.set)

        self.tables_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(5, 0), pady=(0, 5))
        self.tables_scrollbar.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 5), pady=(0, 5))

        # 鼠标滚轮滚动支持（Windows/Linux）
        self._bind_mousewheel(self.tables_canvas)
        self._bind_mousewheel(self.tables_inner)

        # 操作按钮
        action_frame = ttk.Frame(self)
        action_frame.pack(fill=tk.X, padx=5, pady=10)
        self.export_btn = ttk.Button(action_frame, text="4. 导出结构为 Excel", command=self.start_export)
        self.export_btn.pack(side=tk.LEFT, padx=5)

        # 状态
        self.status_var = tk.StringVar(value="准备就绪。")
        status_bar = ttk.Label(self, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)

    # -------------------- UI辅助 --------------------
    def set_status(self, message: str) -> None:
        self.status_var.set(message)
        self.update_idletasks()

    def browse_save_path(self) -> None:
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            title="选择导出文件"
        )
        if path:
            self.save_path_var.set(path)

    # -------------------- 导出主流程 --------------------
    def start_export(self) -> None:
        self.export_btn.config(state=tk.DISABLED)
        self.set_status("开始导出...")
        threading.Thread(target=self._export_impl, daemon=True).start()

    def _export_impl(self) -> None:
        try:
            cfg = self.conn_frame.get_config()
            if not all([cfg["host"], cfg["user"], cfg["database"]]):
                raise ValueError("数据库连接信息不完整！")

            # 保存连接配置以便复用
            self.conn_frame.save_to_config()

            save_path = self.save_path_var.get().strip()
            if not save_path:
                raise ValueError("请先选择导出的 Excel 文件路径！")

            # 若表清单未加载则自动加载一次（默认全选）
            if not self.tables_data:
                self.tables_data = self._read_tables(cfg)
                self._populate_table_vars(self.tables_data, default_selected=True)

            # 获取选中的表
            selected = self._get_selected_tables()
            if not selected:
                raise ValueError("未选择需要导出的表。")

            # 写入Excel
            if self.merge_to_one_sheet.get():
                self._write_excel_merged(save_path, cfg, selected)
            else:
                self._write_excel(save_path, cfg, selected)

            messagebox.showinfo("成功", f"导出完成：\n{save_path}")
            self.set_status("导出完成。")
        except ImportError as e:
            messagebox.showerror("缺少依赖", f"请先安装依赖：openpyxl\n\npip install openpyxl\n\n{e}")
            self.set_status("导出失败：缺少 openpyxl。")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))
            self.set_status(f"导出失败：{e}")
        finally:
            self.export_btn.config(state=tk.NORMAL)

    # -------------------- 数据库读取 --------------------
    def _read_tables(self, db_cfg: Dict[str, object]) -> List[Tuple[str, str]]:
        """
        返回 [(表名, 表注释), ...]
        """
        conn = pymysql.connect(**db_cfg)
        try:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT TABLE_NAME, TABLE_COMMENT FROM INFORMATION_SCHEMA.TABLES "
                    "WHERE TABLE_SCHEMA=%s ORDER BY TABLE_NAME",
                    (db_cfg["database"],)
                )
                return [(r[0], r[1] or "") for r in cursor.fetchall()]
        finally:
            conn.close()

    # -------------------- 表清单 UI --------------------
    def load_tables_list(self) -> None:
        """手动加载/刷新表列表，默认全选"""
        try:
            cfg = self.conn_frame.get_config()
            if not all([cfg["host"], cfg["user"], cfg["database"]]):
                raise ValueError("请先填写完整的数据库连接信息。")
            self.conn_frame.save_to_config()
            self.set_status("正在加载表列表...")
            self.tables_data = self._read_tables(cfg)
            if not self.tables_data:
                raise ValueError("该数据库下没有表。")
            self._populate_table_vars(self.tables_data, default_selected=True)
            self.set_status(f"已加载 {len(self.tables_data)} 张表。")
        except Exception as e:
            messagebox.showerror("加载失败", str(e))
            self.set_status(f"加载表列表失败：{e}")

    def _populate_table_vars(self, tables: List[Tuple[str, str]], default_selected: bool = True) -> None:
        for w in self.tables_inner.winfo_children():
            w.destroy()
        self.table_vars.clear()
        for table_name, comment in tables:
            var = tk.BooleanVar(value=default_selected)
            # 任何勾选变化时刷新按钮文字
            var.trace_add("write", lambda *_: self.update_toggle_button_text())
            self.table_vars[table_name] = var
            text = f"{table_name}" + (f"  （{comment}）" if comment else "")
            ttk.Checkbutton(self.tables_inner, text=text, variable=var).pack(anchor="w", padx=6, pady=2)
        self.update_toggle_button_text()

    def select_all_tables(self) -> None:
        for v in self.table_vars.values():
            v.set(True)
        self.update_toggle_button_text()

    def select_none_tables(self) -> None:
        for v in self.table_vars.values():
            v.set(False)
        self.update_toggle_button_text()

    def toggle_select_tables(self) -> None:
        """一键全选/全不选：若存在未选中则全选，否则全不选"""
        if not self.table_vars:
            return
        any_unchecked = any(not v.get() for v in self.table_vars.values())
        if any_unchecked:
            self.select_all_tables()
        else:
            self.select_none_tables()

    def update_toggle_button_text(self) -> None:
        if not hasattr(self, "toggle_select_btn"):
            return
        if not self.table_vars:
            self.toggle_select_btn.config(text="全选")
            return
        all_checked = all(v.get() for v in self.table_vars.values())
        self.toggle_select_btn.config(text="全不选" if all_checked else "全选")

    def _get_selected_tables(self) -> List[Tuple[str, str]]:
        if not self.table_vars:
            return []
        selected_names = {name for name, v in self.table_vars.items() if v.get()}
        return [(t, c) for t, c in self.tables_data if t in selected_names]

    def _read_columns(self, db_cfg: Dict[str, object], table: str) -> List[Tuple[str, str, str, str, str]]:
        """
        返回列数据：[(name, column_type, is_nullable, is_pk, comment), ...]
        """
        conn = pymysql.connect(**db_cfg)
        try:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT COLUMN_NAME, COLUMN_TYPE, IS_NULLABLE, COLUMN_KEY, COLUMN_COMMENT "
                    "FROM INFORMATION_SCHEMA.COLUMNS "
                    "WHERE TABLE_SCHEMA=%s AND TABLE_NAME=%s "
                    "ORDER BY ORDINAL_POSITION",
                    (db_cfg["database"], table)
                )
                rows = cursor.fetchall()
                result = []
                for name, ctype, nullable, ckey, comment in rows:
                    result.append((
                        name,
                        ctype,
                        "是" if (nullable or "").upper() == "YES" else "否",
                        "是" if (ckey or "").upper() == "PRI" else "否",
                        comment or ""
                    ))
                return result
        finally:
            conn.close()

    # -------------------- 写Excel --------------------
    def _write_excel(self, save_path: str, db_cfg: Dict[str, object], tables: List[Tuple[str, str]]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()

        toc_ws = None
        if self.include_toc.get():
            # 使用默认Sheet作为目录
            toc_ws = wb.active
            toc_ws.title = "目录"
            # 目录头部信息
            toc_ws["A1"] = "库名"
            toc_ws["B1"] = db_cfg.get("database", "")
            toc_ws["A2"] = "服务器"
            toc_ws["B2"] = f"{db_cfg.get('host','')}:{db_cfg.get('port','')}"
            toc_ws["A3"] = "导出时间"
            toc_ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            toc_ws["A4"] = "表数量"
            toc_ws["B4"] = len(tables)
            toc_ws["A1"].font = Font(bold=True)
            toc_ws["A2"].font = Font(bold=True)
            toc_ws["A3"].font = Font(bold=True)
            toc_ws["A4"].font = Font(bold=True)

            # 表头
            toc_headers = ["序号", "表名（点击跳转）", "表描述"]
            for i, h in enumerate(toc_headers, start=1):
                cell = toc_ws.cell(row=6, column=i, value=h)
                cell.font = Font(bold=True)
        else:
            # 删除默认空sheet
            default_ws = wb.active
            wb.remove(default_ws)

        used_names = set()
        toc_row = 7
        idx_no = 1

        for table_name, table_comment in tables:
            columns = self._read_columns(db_cfg, table_name)

            sheet_name = self._safe_sheet_name(table_name)
            # 防重名（可能因裁剪/替换后重复）
            original = sheet_name
            idx = 2
            while sheet_name in used_names:
                suffix = f"_{idx}"
                sheet_name = (original[:31 - len(suffix)] + suffix) if len(original) + len(suffix) > 31 else original + suffix
                idx += 1
            used_names.add(sheet_name)

            ws = wb.create_sheet(title=sheet_name)

            # 顶部信息：表名/表描述
            ws["A1"] = "表名"
            ws["B1"] = table_name
            ws["A2"] = "表描述"
            ws["B2"] = table_comment or ""
            ws["A1"].font = Font(bold=True)
            ws["A2"].font = Font(bold=True)

            # 表头
            headers = ["字段名", "数据类型（长）", "是否为空", "是否为主键", "说明"]
            start_row = 4
            for i, h in enumerate(headers, start=1):
                cell = ws.cell(row=start_row, column=i, value=h)
                cell.font = Font(bold=True)

            # 数据行
            for r_idx, row in enumerate(columns, start=start_row + 1):
                for c_idx, value in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value)

            # 简易列宽
            widths = [16, 20, 10, 12, 50]
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[chr(64 + i)].width = w

            # 目录项
            if toc_ws is not None:
                toc_ws.cell(row=toc_row, column=1, value=idx_no)
                name_cell = toc_ws.cell(row=toc_row, column=2)
                name_cell.value = self._hyperlink_formula(sheet_name, "A1", table_name)
                name_cell.style = "Hyperlink"
                toc_ws.cell(row=toc_row, column=3, value=table_comment or "")
                toc_row += 1
                idx_no += 1

        # 列宽
        if toc_ws is not None:
            toc_ws.column_dimensions["A"].width = 8
            toc_ws.column_dimensions["B"].width = 40
            toc_ws.column_dimensions["C"].width = 60

        wb.save(save_path)

    def _write_excel_merged(self, save_path: str, db_cfg: Dict[str, object], tables: List[Tuple[str, str]]) -> None:
        """将所有表写入一张Sheet，按：表名/表描述/表头/数据 的块顺序排列；可选生成目录"""
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        toc_ws = wb.active
        toc_ws.title = "目录"
        ws = wb.create_sheet(title="数据库结构")

        # 统一列宽
        widths = [16, 20, 10, 12, 50]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w

        current_row = 1
        anchors: Dict[str, int] = {}
        for table_name, table_comment in tables:
            # 顶部信息
            ws.cell(row=current_row, column=1, value="表名").font = Font(bold=True)
            ws.cell(row=current_row, column=2, value=table_name)
            ws.cell(row=current_row + 1, column=1, value="表描述").font = Font(bold=True)
            ws.cell(row=current_row + 1, column=2, value=table_comment or "")

            # 表头
            headers = ["字段名", "数据类型（长）", "是否为空", "是否为主键", "说明"]
            header_row = current_row + 3
            for i, h in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=i, value=h)
                cell.font = Font(bold=True)

            # 数据
            columns = self._read_columns(db_cfg, table_name)
            write_row = header_row + 1
            for row in columns:
                for c_idx, value in enumerate(row, start=1):
                    ws.cell(row=write_row, column=c_idx, value=value)
                write_row += 1

            # 下一个表块，留出一个空行
            current_row = write_row + 1
            anchors[table_name] = header_row  # 记录表在“数据库结构”中的定位行

        # 填充目录或删除目录
        if self.include_toc.get():
            # 目录头部信息
            toc_ws["A1"] = "库名"
            toc_ws["B1"] = db_cfg.get("database", "")
            toc_ws["A2"] = "服务器"
            toc_ws["B2"] = f"{db_cfg.get('host','')}:{db_cfg.get('port','')}"
            toc_ws["A3"] = "导出时间"
            toc_ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            toc_ws["A4"] = "表数量"
            toc_ws["B4"] = len(tables)
            for cell in ("A1","A2","A3","A4"):
                toc_ws[cell].font = Font(bold=True)

            headers = ["序号", "表名（点击跳转）", "表描述"]
            for i, h in enumerate(headers, start=1):
                c = toc_ws.cell(row=6, column=i, value=h)
                c.font = Font(bold=True)
            r = 7
            idx_no = 1
            for table_name, table_comment in tables:
                toc_ws.cell(row=r, column=1, value=idx_no)
                name_cell = toc_ws.cell(row=r, column=2)
                name_cell.value = self._hyperlink_formula("数据库结构", f"A{anchors.get(table_name, 1)}", table_name)
                name_cell.style = "Hyperlink"
                toc_ws.cell(row=r, column=3, value=table_comment or "")
                r += 1
                idx_no += 1
            toc_ws.column_dimensions["A"].width = 8
            toc_ws.column_dimensions["B"].width = 40
            toc_ws.column_dimensions["C"].width = 60
        else:
            # 不需要目录则删除默认sheet
            wb.remove(toc_ws)

        wb.save(save_path)

    def _safe_sheet_name(self, name: str) -> str:
        # Excel工作表名限制：长度<=31，不能包含: \ / ? * [ ]
        invalid = [":", "\\", "/", "?", "*", "[", "]"]
        for ch in invalid:
            name = name.replace(ch, "_")
        name = name.strip()
        return name[:31] if len(name) > 31 else name

    def _hyperlink_formula(self, sheet_name: str, cell_ref: str, display_text: str) -> str:
        """
        构造兼容 Excel/WPS 的 HYPERLINK 公式，支持带空格/中文/特殊字符的工作表名。
        - 工作表名中的单引号需要双写以转义
        - 公式中的双引号需要双写
        """
        sheet_escaped = sheet_name.replace("'", "''")
        text_escaped = (display_text or "").replace('"', '""')
        return f"=HYPERLINK(\"#'{sheet_escaped}'!{cell_ref}\",\"{text_escaped}\")"

    # -------------------- 滚轮支持 --------------------
    def _bind_mousewheel(self, widget: tk.Widget) -> None:
        widget.bind("<Enter>", lambda e: widget.bind_all("<MouseWheel>", self._on_mousewheel_windows), add="+")
        widget.bind("<Leave>", lambda e: widget.unbind_all("<MouseWheel>"), add="+")
        # Linux 支持
        widget.bind("<Button-4>", self._on_mousewheel_linux, add="+")
        widget.bind("<Button-5>", self._on_mousewheel_linux, add="+")

    def _on_mousewheel_windows(self, event) -> None:
        # event.delta 为 120 或 -120 的倍数
        self.tables_canvas.yview_scroll(-1 * int(event.delta / 120), "units")

    def _on_mousewheel_linux(self, event) -> None:
        # Button-4 上滚，Button-5 下滚
        if event.num == 4:
            self.tables_canvas.yview_scroll(-3, "units")
        elif event.num == 5:
            self.tables_canvas.yview_scroll(3, "units")


